"""Ficha maestra de personal por país: plantilla de Excel para carga masiva
+ parseo de lo que suban. Es la fuente de verdad contra la que se cruzan
novedades y liquidaciones (ver `alertas.py`).
"""

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

COLUMNAS = ["identificacion", "nombre", "cargo", "salario_base", "fecha_ingreso", "estado"]


def build_template_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personal"

    headers = ["Identificación", "Nombre completo", "Cargo", "Salario base", "Fecha de ingreso (AAAA-MM-DD)", "Estado (activo/inactivo)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E293B")

    # Fila de ejemplo — no se importa (parse_empleados_excel la reconoce por
    # la identificación literal "EJEMPLO123" y la descarta).
    ws.append(["EJEMPLO123", "Nombre Apellido", "Gerente de marca", 3000000, "2024-01-15", "activo"])
    ws["A2"].font = Font(italic=True, color="94A3B8")
    for col in "BCDEF":
        ws[f"{col}2"].font = Font(italic=True, color="94A3B8")

    ws.append(["Instrucciones: borre la fila de ejemplo (fila 2) y agregue una fila por empleado. "
               "No cambie los encabezados. 'Estado' debe ser 'activo' o 'inactivo'.", None, None, None, None, None])
    ws["A3"].font = Font(italic=True, size=9, color="94A3B8")

    for col, width in zip("ABCDEF", [16, 28, 22, 14, 24, 20]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_empleados_excel(file_stream):
    """Lee un .xlsx subido y devuelve una lista de dicts con las llaves de
    COLUMNAS. Se salta encabezado, filas vacías y la fila de ejemplo de la
    plantilla. No valida exhaustivamente — errores de formato de fecha o
    salario simplemente se guardan como None/texto crudo, para no bloquear
    la carga completa por una fila mal diligenciada."""
    wb = openpyxl.load_workbook(file_stream, data_only=True)
    ws = wb.active
    empleados = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        identificacion = str(row[0]).strip()
        if identificacion.upper() == "EJEMPLO123":
            continue
        nombre = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if not nombre:
            continue
        cargo = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        salario_base = None
        if len(row) > 3 and row[3] not in (None, ""):
            try:
                salario_base = float(row[3])
            except (TypeError, ValueError):
                salario_base = None
        fecha_ingreso = None
        if len(row) > 4 and row[4]:
            val = row[4]
            if isinstance(val, datetime):
                fecha_ingreso = val.date().isoformat()
            else:
                fecha_ingreso = str(val).strip()
        estado = "activo"
        if len(row) > 5 and row[5]:
            estado_raw = str(row[5]).strip().lower()
            if estado_raw in ("inactivo", "activo"):
                estado = estado_raw
        empleados.append({
            "identificacion": identificacion,
            "nombre": nombre,
            "cargo": cargo,
            "salario_base": salario_base,
            "fecha_ingreso": fecha_ingreso,
            "estado": estado,
        })
    return empleados
